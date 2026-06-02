#!/usr/bin/env python3
"""
POC 3 — Stage 3: PPTX Generation using POC_3.pptx template

Uses POC_3.pptx as the visual template and fills in content from the storyboard:

  "1B_Intro-Outro"                      → cover slide (VIT logo / navy bg)
  "1D_Wide Screen"                       → section dividers (blurred corridor)
  "60:40 Icon and text box with bullets" → bullets / summary slides
  "2C_Image Dark + Text (right, bottom)" → visual / quote slides

Generates one Gemini Imagen image per content slide and embeds them.

Outputs  →  <output_dir>/storyboard.pptx
"""

import sys
import json
import os
from copy import deepcopy
from pathlib import Path

from lxml import etree as _etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.oxml.ns import qn

# ── Template path ──────────────────────────────────────────────────────────────
_DEFAULT_TEMPLATE = str(Path(__file__).parent.parent / "POC_3.pptx")

# ── Layout names (must match names in POC_3.pptx) ─────────────────────────────
L_INTRO      = "1B_Intro-Outro"           # VIT logo / navy bg — cover & outro
L_VID_TITLE  = "2_1C_Video Title"         # campus photo + dark overlay + topic
L_WIDE       = "1D_Wide Screen"           # blurred corridor — section dividers
L_BULLETS    = "60:40 Icon and text box with bullets"
L_IMAGE_DARK = "2C_Image Dark + Text (right, bottom)"

# ── Placeholder indices ────────────────────────────────────────────────────────
PH_TITLE     = 0    # slide title          (L_BULLETS)
PH_MOD_TITLE = 34   # top banner           (L_BULLETS)
PH_BODY      = 41   # bullets body         (L_BULLETS)
PH_ICON      = 43   # small icon square    (L_BULLETS)
PH_IMG_FULL  = 23   # full-bleed image     (L_IMAGE_DARK)
PH_CAPTION   = 22   # caption / text box   (L_IMAGE_DARK)
PH_VID_BODY  = 12   # topic title body     (L_VID_TITLE)

# ── Right-panel image area for bullets layout ──────────────────────────────────
# The right 40% of a 13.33" slide starts at ~8.1"
_RP_LEFT  = Inches(8.2)
_RP_W     = Inches(4.9)
_RP_IMG_H = Emu(_RP_W * 9 / 16)                  # 16:9 → ≈ 2.76"
_RP_TOP   = Emu((Inches(7.5) - _RP_IMG_H) / 2)   # vertically centred

# ── Relationship namespace for r:embed remapping ──────────────────────────────
_NS_R      = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
_EMBED_ATTR = f'{{{_NS_R}}}embed'


# ── Helpers ────────────────────────────────────────────────────────────────────

def _find_layout(prs, name):
    """Find a slide layout by name across all masters; fall back to first."""
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if layout.name == name:
                return layout
    return prs.slide_masters[0].slide_layouts[0]


def _clear_slides(prs):
    """Remove every existing slide from the presentation (keeps masters/layouts)."""
    sldIdLst = prs.slides._sldIdLst
    rids = [el.get(qn("r:id")) for el in list(sldIdLst)]
    for rId in rids:
        if rId:
            try:
                prs.part.drop_rel(rId)
            except Exception:
                pass
    for child in list(sldIdLst):
        sldIdLst.remove(child)


def _ph(slide, idx):
    """Return placeholder by format-idx, or None."""
    for p in slide.placeholders:
        if p.placeholder_format.idx == idx:
            return p
    return None


def _fill_ph(slide, idx, text, size=None, bold=None, color=None, italic=None,
             align=None):
    """Fill a placeholder with a single run of text."""
    p = _ph(slide, idx)
    if p is None:
        return None
    tf = p.text_frame
    tf.clear()
    tf.word_wrap = True
    para = tf.paragraphs[0]
    if align:
        para.alignment = align
    run = para.add_run()
    run.text = text
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = color
    return p


def _fill_bullets(slide, idx, bullets):
    """Fill a placeholder with one paragraph per bullet string.

    Rebuilds the txBody paragraphs at XML level to prevent LibreOffice
    from bleeding the layout's inherited paragraph content as ghost text.
    Space-before is scaled dynamically so fewer bullets breathe more.
    """
    ph = _ph(slide, idx)
    if ph is None:
        return None

    txBody = ph._element.find(qn('p:txBody'))
    if txBody is None:
        return None

    # Nuke ALL layout-inherited <a:p> elements — this is what causes ghost text
    for ap in list(txBody.findall(qn('a:p'))):
        txBody.remove(ap)

    # Enable normAutofit so text shrinks rather than overflows the fixed box
    bodyPr = txBody.find(qn('a:bodyPr'))
    if bodyPr is not None:
        for tag in (qn('a:normAutofit'), qn('a:spAutoFit'), qn('a:noAutofit')):
            el = bodyPr.find(tag)
            if el is not None:
                bodyPr.remove(el)
        _etree.SubElement(bodyPr, qn('a:normAutofit'))

    # spcBef in hundredths-of-a-point: fewer bullets → more breathing room
    n = max(1, len(bullets))
    spc = {1: 2400, 2: 1800, 3: 1200, 4: 800, 5: 400}.get(n, 200)

    for b in bullets:
        ap   = _etree.SubElement(txBody, qn('a:p'))
        pPr  = _etree.SubElement(ap, qn('a:pPr'))
        pPr.set('lvl', '0')
        spcBef = _etree.SubElement(pPr, qn('a:spcBef'))
        spcPts = _etree.SubElement(spcBef, qn('a:spcPts'))
        spcPts.set('val', str(spc))
        r    = _etree.SubElement(ap, qn('a:r'))
        rPr  = _etree.SubElement(r, qn('a:rPr'))
        rPr.set('lang', 'en-US')
        rPr.set('dirty', '0')
        t    = _etree.SubElement(r, qn('a:t'))
        t.text = b

    ph.text_frame.word_wrap = True
    return ph


def _textbox(slide, text, left, top, width, height,
             size=18, bold=False, italic=False,
             color=RGBColor(0xFF, 0xFF, 0xFF), align=PP_ALIGN.LEFT):
    tb  = slide.shapes.add_textbox(left, top, width, height)
    tf  = tb.text_frame
    tf.word_wrap = True
    para = tf.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.size   = Pt(size)
    run.font.bold   = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return tb


def _insert_picture_ph(slide, idx, img_path):
    """Insert image into a PICTURE-type placeholder; fall back to add_picture."""
    p = _ph(slide, idx)
    if p is None:
        return None
    try:
        p.insert_picture(str(img_path))
        return p
    except Exception:
        # Fallback: place image at the placeholder's position/size
        slide.shapes.add_picture(str(img_path), p.left, p.top, p.width, p.height)
        return None


def _write_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text or ""


def _clean_layout_prompts(prs):
    """Wipe prompt/hint paragraphs from the bullets layout's body placeholder.

    LibreOffice merges the LAYOUT's <a:p> nodes with the slide's <a:p> nodes
    during rendering, causing "Click to edit text" and empty paragraphs to
    appear ghost-style between real bullets. Clearing the layout paragraphs
    once (after loading the template, before building slides) eliminates this.
    """
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if layout.name != L_BULLETS:
                continue
            for ph in layout.placeholders:
                if ph.placeholder_format.idx != PH_BODY:
                    continue
                txBody = ph._element.find(qn('p:txBody'))
                if txBody is None:
                    continue
                for ap in list(txBody.findall(qn('a:p'))):
                    txBody.remove(ap)
                # OOXML requires at least one <a:p>; leave a clean empty one
                ap = _etree.SubElement(txBody, qn('a:p'))
                _etree.SubElement(ap, qn('a:endParaRPr')).set('lang', 'en-US')


def _borrow_bg_picture(slide, layout):
    """
    Copy the full-bleed background picture from a layout's spTree into the
    slide's spTree so LibreOffice renders it (layouts marked userDrawn='1'
    are not automatically inherited by derived slides in LibreOffice).
    Only copies pictures that span at least 10 inches wide (full-bleed).
    """
    layout_spTree = layout._element.find('.//' + qn('p:spTree'))
    if layout_spTree is None:
        return
    slide_spTree = slide._element.find('.//' + qn('p:spTree'))

    insert_idx = 2  # after nvGrpSpPr[0] and grpSpPr[1], before any text boxes
    for elem in layout_spTree:
        if elem.tag.split('}')[-1] != 'pic':
            continue
        xfrm = elem.find('.//' + qn('a:xfrm'))
        ext  = xfrm.find(qn('a:ext')) if xfrm is not None else None
        if ext is None or int(ext.get('cx', '0')) < Inches(10):
            continue  # skip non-full-bleed pictures (person photos, icons etc.)

        pic_copy = deepcopy(elem)
        # Remap every r:embed reference from layout's rIds to new slide rIds
        for el in pic_copy.iter():
            old_rId = el.get(_EMBED_ATTR)
            if old_rId:
                try:
                    rel = layout.part.rels[old_rId]
                    new_rId = slide.part.relate_to(rel.target_part, rel.reltype)
                    el.set(_EMBED_ATTR, new_rId)
                except (KeyError, AttributeError):
                    pass
        slide_spTree.insert(insert_idx, pic_copy)
        insert_idx += 1
        break  # one full-bleed background is enough


# ── Slide builders ─────────────────────────────────────────────────────────────

def _add_cover(prs, doc_title):
    """Intro card: VIT logo on navy background + module title in a bottom bar."""
    layout = _find_layout(prs, L_INTRO)
    slide  = prs.slides.add_slide(layout)

    # Solid footer strip below the VIT crest (crest ribbon ends ~6.2")
    strip = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0), Inches(6.1), Inches(13.33), Inches(1.4),
    )
    strip.fill.solid()
    strip.fill.fore_color.rgb = RGBColor(0x0A, 0x14, 0x28)
    strip.line.fill.background()

    _textbox(slide, doc_title,
             Inches(0.6), Inches(6.2), Inches(12.0), Inches(1.2),
             size=26, bold=True,
             color=RGBColor(0xFF, 0xFF, 0xFF))
    return slide


def _add_section_divider(prs, heading, sec_num):
    """Blurred-corridor photo background + dark navy text band with section heading."""
    layout = _find_layout(prs, L_WIDE)
    slide  = prs.slides.add_slide(layout)

    # Inject the corridor photo directly so LibreOffice renders it
    _borrow_bg_picture(slide, layout)

    # Solid dark-navy band across the full width (centre vertically)
    band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0), Inches(2.0), Inches(13.33), Inches(3.2),
    )
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(0x0D, 0x1B, 0x2A)
    band.line.fill.background()

    # Section label in template gold/orange (#FFC000)
    _textbox(slide, f"SECTION  {sec_num}",
             Inches(0.8), Inches(2.25), Inches(11.0), Inches(0.5),
             size=14, bold=True,
             color=RGBColor(0xFF, 0xC0, 0x00))

    # Section heading in white
    _textbox(slide, heading,
             Inches(0.8), Inches(2.85), Inches(11.0), Inches(2.1),
             size=40, bold=True,
             color=RGBColor(0xFF, 0xFF, 0xFF))
    return slide


def _build_bullets_slide(prs, sd, img_path, doc_title=""):
    """
    60:40 layout: left panel has module title banner, slide title, bullets.
    Right panel: Imagen image positioned over the hexagonal pattern.
    """
    layout = _find_layout(prs, L_BULLETS)
    slide  = prs.slides.add_slide(layout)

    # Top banner and title
    _fill_ph(slide, PH_MOD_TITLE, doc_title, size=11)
    # Override italic/color from placeholder default (template default is italic teal)
    _fill_ph(slide, PH_TITLE, sd.get("title", ""),
             bold=True, italic=False, color=RGBColor(0x1A, 0x1A, 0x2E))

    # Bullets body
    bullets = sd.get("bullets", [])
    _fill_bullets(slide, PH_BODY, bullets)

    # Remove the icon placeholder (we have no icon; avoids empty box artefact)
    icon_ph = _ph(slide, PH_ICON)
    if icon_ph is not None:
        icon_ph._element.getparent().remove(icon_ph._element)

    # Imagen image in the right panel (overlaid on hexagonal background)
    if img_path and Path(img_path).exists():
        slide.shapes.add_picture(str(img_path), _RP_LEFT, _RP_TOP, _RP_W, _RP_IMG_H)

    _write_notes(slide, sd.get("speaker_notes", ""))
    return slide


def _build_visual_slide(prs, sd, img_path, doc_title=""):
    """Full-bleed Imagen image + slide title in the caption box (bottom-right)."""
    layout = _find_layout(prs, L_IMAGE_DARK)
    slide  = prs.slides.add_slide(layout)

    if img_path and Path(img_path).exists():
        _insert_picture_ph(slide, PH_IMG_FULL, img_path)

    _fill_ph(slide, PH_CAPTION, sd.get("title", ""),
             size=18, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    _write_notes(slide, sd.get("speaker_notes", ""))
    return slide


def _build_quote_slide(prs, sd, img_path, doc_title=""):
    """Full-bleed image + quote text in the caption box."""
    layout = _find_layout(prs, L_IMAGE_DARK)
    slide  = prs.slides.add_slide(layout)

    if img_path and Path(img_path).exists():
        _insert_picture_ph(slide, PH_IMG_FULL, img_path)

    bullets   = sd.get("bullets", [])
    quote_txt = f"“{bullets[0]}”" if bullets else sd.get("title", "")
    _fill_ph(slide, PH_CAPTION, quote_txt,
             size=16, italic=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    _write_notes(slide, sd.get("speaker_notes", ""))
    return slide


def _build_summary_slide(prs, sd, img_path, doc_title=""):
    sd2 = dict(sd)
    sd2.setdefault("title", "Key Takeaways")
    return _build_bullets_slide(prs, sd2, img_path, doc_title)


_BUILDERS = {
    "bullets": _build_bullets_slide,
    "visual":  _build_visual_slide,
    "quote":   _build_quote_slide,
    "table":   _build_bullets_slide,
    "summary": _build_summary_slide,
}


# ── Image generation ───────────────────────────────────────────────────────────

def _generate_images(storyboard, images_dir, gemini_api_key):
    from poc3_image_gen import generate_slide_image
    paths     = {}
    doc_title = storyboard.get("title", "Course Module")

    cover_prompt = (
        f"Professional educational course cover visual representing '{doc_title}'. "
        "Abstract geometric shapes with deep blue and gold palette, modern flat design, "
        "no text or lettering anywhere in the image."
    )
    print("          [cover]  Generating …", end=" ", flush=True)
    p = generate_slide_image(cover_prompt,
                             str(images_dir / "cover.png"),
                             gemini_api_key)
    paths["cover"] = p
    print("done" if p else "failed")

    for sec in storyboard.get("sections", []):
        for slide in sec.get("slides", []):
            num    = slide["slide_number"]
            prompt = slide.get("image_prompt", "") or (
                f"Professional educational illustration for a presentation slide titled "
                f"'{slide.get('title', '')}'. Flat design, modern palette, no text."
            )
            out = str(images_dir / f"slide_{num:03d}.png")
            print(f"          [slide {num:2d}]  Generating …", end=" ", flush=True)
            p = generate_slide_image(prompt, out, gemini_api_key)
            paths[num] = p
            print("done" if p else "failed")

    return paths


# ── Stage 3 runner ─────────────────────────────────────────────────────────────

def run_stage3(storyboard_path, output_dir, gemini_api_key=None,
               template_path=None):
    with open(storyboard_path, encoding="utf-8") as f:
        storyboard = json.load(f)

    template_path = template_path or _DEFAULT_TEMPLATE

    out        = Path(output_dir)
    images_dir = out / "images"
    images_dir.mkdir(parents=True, exist_ok=True)

    doc_title = storyboard.get("title", "Untitled")
    sections  = storyboard.get("sections", [])

    # ── Image generation ──────────────────────────────────────────────────────
    img_paths = {}
    if gemini_api_key:
        print(f"[Stage 3/3] Generating images with Gemini Imagen …")
        img_paths = _generate_images(storyboard, images_dir, gemini_api_key)
    else:
        print("[Stage 3/3] Skipping image generation (no GEMINI_API_KEY set)")
        # Reuse any images already present in the output directory
        cover_img = images_dir / "cover.png"
        if cover_img.exists():
            img_paths["cover"] = str(cover_img)
        for sec in storyboard.get("sections", []):
            for sd in sec.get("slides", []):
                num = sd["slide_number"]
                p = images_dir / f"slide_{num:03d}.png"
                if p.exists():
                    img_paths[num] = str(p)
        if img_paths:
            print(f"[Stage 3/3] Reusing {len(img_paths)} existing image(s) from {images_dir.name}/")

    # ── Load template & clear existing slides ─────────────────────────────────
    print(f"[Stage 3/3] Building PPTX from template: {Path(template_path).name} …")
    prs = Presentation(template_path)
    _clear_slides(prs)
    _clean_layout_prompts(prs)   # prevent LibreOffice from ghosting layout text

    # ── Cover ─────────────────────────────────────────────────────────────────
    _add_cover(prs, doc_title)
    slide_count = 1

    # ── Sections ──────────────────────────────────────────────────────────────
    for sec in sections:
        heading = sec.get("section_heading", "")
        sec_id  = sec.get("section_id", 0)
        slides  = sec.get("slides", [])

        _add_section_divider(prs, heading, sec_id)
        slide_count += 1

        for sd in slides:
            num     = sd["slide_number"]
            scene   = sd.get("scene_type", "bullets")
            builder = _BUILDERS.get(scene, _build_bullets_slide)
            builder(prs, sd, img_paths.get(num), doc_title)
            slide_count += 1
            print(f"            [{scene:8s}] slide {num:2d} — {sd.get('title','')[:55]}")

    # ── Save ──────────────────────────────────────────────────────────────────
    pptx_path = out / "storyboard.pptx"
    prs.save(str(pptx_path))
    size_kb = pptx_path.stat().st_size / 1024
    print(f"\n[Stage 3 Complete] {slide_count} slides → {pptx_path}  ({size_kb:.0f} KB)")

    manifest = {"pptx": str(pptx_path), "total_slides": slide_count}
    with open(out / "stage3_manifest.json", "w") as f:
        json.dump(manifest, f, indent=2)
    return manifest


# ── CLI ────────────────────────────────────────────────────────────────────────

def _load_env():
    env = Path(__file__).parent.parent / ".env"
    if env.exists():
        with open(env) as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, _, v = line.partition("=")
                    os.environ.setdefault(k.strip(), v.strip())


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    _load_env()
    gemini_key    = os.environ.get("GEMINI_API_KEY")
    template_path = sys.argv[3] if len(sys.argv) > 3 else None
    result = run_stage3(sys.argv[1], sys.argv[2],
                        gemini_api_key=gemini_key,
                        template_path=template_path)
    print()
    for k, v in result.items():
        print(f"  {k}: {v}")
